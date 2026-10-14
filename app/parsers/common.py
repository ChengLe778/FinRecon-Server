# -*- coding: utf-8 -*-
"""解析层公共工具：文件读取（xls/xlsx 统一）与表头映射。"""
from __future__ import annotations

from typing import Any, Dict, Iterable, Iterator, List, Optional, Sequence
import os

import openpyxl
import xlrd

from .base import ParseError


def _is_xls(path: str) -> bool:
    return os.path.splitext(path)[1].lower() == ".xls"


def read_rows(path: str, sheet: Optional[str] = None,
              sheet_index: int = 0) -> List[List[Any]]:
    """读取整个 sheet 为二维数组。xls 用 xlrd，xlsx 用 openpyxl（只读模式，省内存）。"""
    if not os.path.exists(path):
        raise ParseError(f"文件不存在：{path}")
    if _is_xls(path):
        book = xlrd.open_workbook(path)
        if sheet is not None:
            if sheet not in book.sheet_names():
                raise ParseError(f"文件缺少工作表「{sheet}」，实际为 {book.sheet_names()}")
            sh = book.sheet_by_name(sheet)
        else:
            sh = book.sheet_by_index(sheet_index)
        return [[sh.cell_value(r, c) for c in range(sh.ncols)] for r in range(sh.nrows)]

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet is not None:
            if sheet not in wb.sheetnames:
                raise ParseError(f"文件缺少工作表「{sheet}」，实际为 {wb.sheetnames}")
            ws = wb[sheet]
        else:
            ws = wb.worksheets[sheet_index]
        return [list(row) for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def sheet_names(path: str) -> List[str]:
    if _is_xls(path):
        return xlrd.open_workbook(path).sheet_names()
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def build_header_map(header: Sequence[Any], columns: Dict[str, List[str]]) -> Dict[str, int]:
    """把表头行 + 列配置 → {字段名: 列下标}。

    匹配优先级：精确相等 → 前缀匹配 → 包含匹配。
    中行列名为「中文[ English ]」形式，靠前缀匹配即可命中。
    """
    heads = [("" if h is None else str(h).strip()) for h in header]
    result: Dict[str, int] = {}
    for field, candidates in columns.items():
        idx = None
        for cand in candidates:                               # 1) 精确
            for i, h in enumerate(heads):
                if h == cand:
                    idx = i
                    break
            if idx is not None:
                break
        if idx is None:                                        # 2) 前缀
            for cand in candidates:
                for i, h in enumerate(heads):
                    if h.startswith(cand):
                        idx = i
                        break
                if idx is not None:
                    break
        if idx is None:                                        # 3) 包含
            for cand in candidates:
                for i, h in enumerate(heads):
                    if cand and cand in h:
                        idx = i
                        break
                if idx is not None:
                    break
        if idx is not None:
            result[field] = idx
    return result


def cell(row: Sequence[Any], idx_map: Dict[str, int], field: str,
         default: Any = None) -> Any:
    """按字段名安全取单元格。"""
    i = idx_map.get(field)
    if i is None or i >= len(row):
        return default
    return row[i]


def require_fields(idx_map: Dict[str, int], fields: Iterable[str], parser_name: str):
    missing = [f for f in fields if f not in idx_map]
    if missing:
        raise ParseError(
            f"{parser_name}：未能在源文件中找到必需的列 {missing}。"
            f"已识别到的列：{sorted(idx_map.keys())}"
        )
