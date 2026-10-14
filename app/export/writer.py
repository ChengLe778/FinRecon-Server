# -*- coding: utf-8 -*-
"""导出工具：openpyxl 生成多 sheet 的 xlsx，标准库 csv 生成带 BOM 的 CSV。"""
from __future__ import annotations

import io
from typing import Any, Dict, List, Sequence

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

MONEY_FMT = "#,##0.00"
HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
HEADER_FONT = Font(bold=True, size=11)
TOTAL_FONT = Font(bold=True)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _width_of(v: Any) -> int:
    """按内容估算显示宽度：中文按 2 个字符计。"""
    s = "" if v is None else str(v)
    w = 0
    for ch in s:
        w += 2 if ord(ch) > 127 else 1
    return w


def sheet_from_rows(ws, headers: Sequence[str], rows: Sequence[Sequence[Any]],
                    money_cols: Sequence[int] = (), total_row: Sequence[Any] = None):
    """把一个二维表写进 worksheet：表头加粗、金额列数字格式、列宽自适应、冻结首行。"""
    ws.append(list(headers))
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER

    for r in rows:
        ws.append(list(r))

    if total_row:
        ws.append(list(total_row))
        tr = ws.max_row
        for c in range(1, len(headers) + 1):
            ws.cell(row=tr, column=c).font = TOTAL_FONT

    # 金额列：右对齐 + 两位小数
    for c in money_cols:
        for r in range(2, ws.max_row + 1):
            cell = ws.cell(row=r, column=c)
            if isinstance(cell.value, (int, float)):
                cell.number_format = MONEY_FMT
            cell.alignment = Alignment(horizontal="right")

    # 列宽自适应（按表头与内容的最长值估算，上限 46）
    for i, h in enumerate(headers, start=1):
        w = _width_of(h)
        for r in range(2, min(ws.max_row, 500) + 1):
            w = max(w, _width_of(ws.cell(row=r, column=i).value))
        ws.column_dimensions[get_column_letter(i)].width = min(max(w + 4, 10), 46)

    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"


def build_xlsx(sheets: List[Dict[str, Any]]) -> io.BytesIO:
    """sheets: [{name, headers, rows, money_cols, total_row}] → xlsx 字节流"""
    wb = Workbook()
    wb.remove(wb.active)
    for sh in sheets:
        ws = wb.create_sheet(title=sh["name"][:31])
        sheet_from_rows(ws, sh["headers"], sh.get("rows", []),
                        sh.get("money_cols", ()), sh.get("total_row"))
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def build_csv(headers: Sequence[str], rows: Sequence[Sequence[Any]]) -> io.BytesIO:
    """CSV（utf-8-sig，保证 Excel 打开中文不乱码）"""
    import csv
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(list(headers))
    for r in rows:
        w.writerow(list(r))
    raw = buf.getvalue().encode("utf-8-sig")
    out = io.BytesIO(raw)
    out.seek(0)
    return out
